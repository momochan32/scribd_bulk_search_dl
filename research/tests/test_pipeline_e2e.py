"""Uji ujung-ke-ujung pada PDF sintetis: text layer, halaman gambar (OCR), duplikat, laporan PDF & Excel."""

import shutil
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from solcoat_research.__main__ import main
from solcoat_research.extract import DEFAULT_TESSDATA, OcrSettings
from solcoat_research.pipeline import ScanConfig, run_scan
from solcoat_research.store import search

TEXT_PAGE = """Laporan Kerja Praktek PT Pupuk Kalimantan Timur
Unit ammonia Kaltim-3
Primary Reformer (1-H-101)
Fungsi : tempat reaksi steam reforming
Temperatur outlet : 800 °C
Tekanan : 38 kg/cm2
Jumlah tube : 168
Program efisiensi energi dijalankan di pabrik amoniak."""

OCR_PAGE = "Auxiliary Boiler (B-2201)\nTekanan : 45 kg/cm2\nTemperatur : 510 C"

HAS_OCR = shutil.which("tesseract") is not None and (Path(DEFAULT_TESSDATA) / "ind.traineddata").exists()


def _text_pdf(path: Path, body: str) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), body, fontsize=11)
    doc.save(path)


def _image_only_pdf(path: Path, body: str) -> None:
    src = pymupdf.open()
    page = src.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 400), body, fontsize=20)
    pixmap = page.get_pixmap(dpi=200)
    out = pymupdf.open()
    target = out.new_page()
    target.insert_image(target.rect, pixmap=pixmap)
    out.save(path)


@pytest.fixture
def corpus(tmp_path: Path) -> Path:
    root = tmp_path / "Momo_Rescribd"
    (root / "Pupuk_Kaltim").mkdir(parents=True)
    (root / "Tes baru").mkdir()
    _text_pdf(root / "Pupuk_Kaltim" / "Laporan-Kerja-Praktek-PKT_111111.pdf", TEXT_PAGE)
    shutil.copy(root / "Pupuk_Kaltim" / "Laporan-Kerja-Praktek-PKT_111111.pdf", root / "Tes baru" / "salinan.pdf")
    _image_only_pdf(root / "Pupuk_Kaltim" / "scan-boiler_222222.pdf", OCR_PAGE)
    (root / ".hidden").mkdir()
    _text_pdf(root / ".hidden" / "abaikan.pdf", "tidak dibaca")
    return root


def test_scan_extracts_relations_and_detects_duplicates(corpus: Path, tmp_path: Path):
    result = run_scan(ScanConfig(corpus, tmp_path / "out", workers=2, ocr=OcrSettings(enabled=HAS_OCR)),
                      log=lambda _: None)

    assert [d.name for d in result.duplicates] == ["salinan.pdf"]
    reformer = next(e for e in result.equipment if e.key == "primary_reformer")
    assert (reformer.company, reformer.plant, reformer.tag) == ("Pupuk Kalimantan Timur", "Kaltim 3", "1-H-101")
    params = {r.record.fact.param_key: r.record.fact for r in result.facts if r.equipment_id == reformer.id}
    assert params["suhu_outlet"].value_std == 800 and params["suhu_outlet"].confidence == "tinggi"
    assert params["jumlah_tube"].value == 168
    assert search(result.db_path, '"primary reformer"')


@pytest.mark.skipif(not HAS_OCR, reason="Tesseract + tessdata ind belum terpasang (jalankan setup_ocr.sh)")
def test_image_only_page_is_read_with_ocr(corpus: Path, tmp_path: Path):
    result = run_scan(ScanConfig(corpus, tmp_path / "out", workers=2), log=lambda _: None)
    scan_doc = next(a for a in result.documents if a.name.startswith("scan-boiler"))
    assert scan_doc.ocr_pages == 1
    boiler_facts = [r.record.fact for r in result.facts if r.record.fact.equipment
                    and r.record.fact.equipment.key == "aux_boiler"]
    assert any(f.param_key == "tekanan" and f.value == 45 for f in boiler_facts)
    assert all(any(fl.startswith("angka hasil OCR") for fl in f.flags) for f in boiler_facts if f.value is not None)


def test_cli_writes_pdf_and_excel(corpus: Path, tmp_path: Path, capsys):
    out = tmp_path / "cli_out"
    assert main(["scan", str(corpus), "--out", str(out), "--workers", "2", "--no-ocr"]) == 0
    pdf = next(out.glob("Laporan_Riset_Solcoat_*.pdf"))
    xlsx = next(out.glob("Data_Riset_Solcoat_*.xlsx"))
    text = "".join(page.get_text() for page in pymupdf.open(pdf))
    assert "Laporan Riset Sumber PDF" in text and "Primary Reformer" in text and "1-H-101" in text
    assert {"Dokumen", "Peralatan", "Fakta", "Perlu_Verifikasi", "Konteks_Strategis", "Duplikat"} <= set(
        load_workbook(xlsx).sheetnames)
    assert main(["search", str(out), "reformer"]) == 0
    assert "h.1" in capsys.readouterr().out


def test_cli_reports_missing_folder(tmp_path: Path):
    assert main(["scan", str(tmp_path / "tidak-ada")]) == 2
    assert main(["search", str(tmp_path), "apa"]) == 2
