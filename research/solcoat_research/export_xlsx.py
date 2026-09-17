"""Ekspor Excel lengkap untuk analis: semua fakta (termasuk keyakinan rendah) bisa difilter."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import report_data as rd
from .pipeline import ScanResult

HEADER_FILL = PatternFill("solid", fgColor="1F5A20")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MAX_COL_WIDTH = 60


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, start=1):
        sample = [len(str(r[idx - 1])) for r in rows[:200] if r[idx - 1] is not None]
        ws.column_dimensions[get_column_letter(idx)].width = min(max([len(header)] + sample) + 2, MAX_COL_WIDTH)


def _fact_rows(result: ScanResult) -> list[list]:
    rows = []
    for r in result.facts:
        f = r.record.fact
        eq = f.equipment
        rows.append([
            r.row_id, r.record.company, f.plant, eq.label if eq else None, eq.tag if eq else None,
            eq.tier if eq else None, eq.component if eq else None, f.category, f.param_label, f.raw, f.value,
            f.value_max, f.unit, f.value_std, f.value_max_std, f.std_unit, f.text_value, f.qualifier, f.confidence,
            f.method, r.record.page_method, r.doc_name, f.page_no, r.source_grade, "; ".join(f.flags),
            r.verification.status if r.verification else "belum diverifikasi",
            r.verification.label if r.verification else None,
            r.verification.corrected_value if r.verification else None,
            r.verification.note if r.verification else None, f.snippet,
        ])
    return rows


def write_xlsx(result: ScanResult, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    duplicates = {d.path: d.original for d in result.duplicates}
    _sheet(wb, "Dokumen", ["Dokumen", "Topik", "Perusahaan", "Jenis sumber", "Grade", "Halaman", "Halaman OCR",
                           "Halaman gagal", "Bahasa", "Relevansi", "Skor", "Duplikat dari", "Error", "Path"],
           [[a.name, a.topic, a.company, a.source_type, a.source_grade, a.page_count, a.ocr_pages, a.failed_pages,
             a.lang, a.relevance, a.relevance_score, duplicates.get(a.path), a.error, a.path]
            for a in sorted(result.documents, key=lambda a: -a.relevance_score)])
    views = rd.target_equipment(result)
    summaries = {v.record.id: rd.key_summary(v) for v in views}
    _sheet(wb, "Peralatan", ["Perusahaan", "Pabrik", "Peralatan", "Tag", "Tier", "Jumlah fakta", "Dokumen",
                             "Di laporan PDF", "Ringkasan"],
           [[e.company, e.plant, e.label, e.tag, e.tier, e.fact_count, "; ".join(e.doc_names),
             "ya" if e.id in summaries else "tidak", summaries.get(e.id, "")] for e in result.equipment])
    _sheet(wb, "Fakta", ["ID", "Perusahaan", "Pabrik", "Peralatan", "Tag", "Tier", "Komponen", "Kategori",
                         "Parameter", "Nilai asli", "Nilai", "Nilai maks", "Satuan asli", "Nilai baku",
                         "Nilai baku maks", "Satuan baku", "Teks", "Kualifier", "Keyakinan", "Metode relasi",
                         "Metode halaman", "Dokumen", "Halaman", "Grade sumber", "Catatan", "Status", "Label",
                         "Koreksi (satuan baku)", "Catatan verifikasi", "Kutipan"],
           _fact_rows(result))
    _sheet(wb, "Perlu_Verifikasi", ["Peralatan", "Parameter", "Nilai asli", "Masalah", "Dokumen", "Halaman"],
           [[(r.record.fact.equipment.label if r.record.fact.equipment else ""), r.record.fact.param_label,
             r.record.fact.raw, "; ".join(rd.display_flags(r.record.fact, include_ocr=True)), r.doc_name,
             r.record.fact.page_no] for r in result.facts if r.record.fact.flags and r.verification is None])
    _sheet(wb, "Konteks_Strategis", ["Dokumen", "Halaman", "Kata kunci", "Kalimat"],
           [[a.name, pg, kw, s] for a in rd.active_documents(result) for pg, kw, s in a.strategic])
    _sheet(wb, "Duplikat", ["Dokumen", "Duplikat dari", "Alasan"],
           [[d.name, d.original, d.reason] for d in result.duplicates])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
